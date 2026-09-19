"""Local extraction fallbacks for sources that Docling cannot process.

These adapters intentionally return raw extraction contracts. Source parsers
remain responsible for canonical-node construction and provenance assignment.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover - exercised when PyMuPDF is not installed.
    fitz = None

try:
    import pdfplumber
except ImportError:  # pragma: no cover - exercised when pdfplumber is not installed.
    pdfplumber = None

from backend.ingestion.extraction.contracts import ExtractedPage
from backend.ingestion.extraction.contracts import ExtractionResult
from backend.ingestion.utils import normalize_text


class PdfFallbackExtractor:
    """Extracts PDF text, coordinates, image metadata, and tables without Docling."""

    # This function uses PyMuPDF first and retains pypdf as a minimal dependency fallback.
    def extract(self, source_path: str | Path) -> ExtractionResult:
        if fitz is not None:
            try:
                return self._extract_with_pymupdf(source_path)
            except Exception as error:
                result = self._extract_with_pypdf(source_path)
                result.warnings.append(f"PyMuPDF failed: {type(error).__name__}; pypdf recovery used.")
                return result
        return self._extract_with_pypdf(source_path)

    # This function extracts ordered text blocks and image metadata with PyMuPDF.
    def _extract_with_pymupdf(self, source_path: str | Path) -> ExtractionResult:
        pages: list[ExtractedPage] = []
        warnings: list[str] = []
        with fitz.open(str(source_path)) as document:
            for page_number, page in enumerate(document, start=1):
                raw_blocks = page.get_text("blocks", sort=True)
                text_blocks = [
                    {
                        "text": normalize_text(block[4]) or "",
                        "bbox": {"x0": block[0], "y0": block[1], "x1": block[2], "y1": block[3]},
                        "block_number": block[5],
                    }
                    for block in raw_blocks
                    if len(block) >= 7 and block[6] == 0 and normalize_text(block[4])
                ]
                text = "\n".join(block["text"] for block in text_blocks)
                images = [
                    {"xref": image[0], "width": image[2], "height": image[3], "extension": image[7]}
                    for image in page.get_images(full=True)
                ]
                try:
                    tables = self._extract_page_tables(source_path, page_number)
                except Exception as error:
                    tables = []
                    warnings.append(f"Table extraction failed on page {page_number}: {type(error).__name__}")
                if not text:
                    warnings.append(f"Low-text or image-only PDF page detected: {page_number}")
                pages.append(ExtractedPage(number=page_number, text=text, text_blocks=text_blocks, tables=tables, images=images))
        return ExtractionResult(extractor_name="pymupdf", extraction_mode="fallback_library", pages=pages, warnings=warnings)

    # This function uses pdfplumber table extraction only when the optional package is installed.
    def _extract_page_tables(self, source_path: str | Path, page_number: int) -> list[list[list[str]]]:
        if pdfplumber is None:
            return []
        with pdfplumber.open(str(source_path)) as document:
            page = document.pages[page_number - 1]
            return [
                [[normalize_text(str(cell or "")) or "" for cell in row] for row in table]
                for table in page.extract_tables()
                if table
            ]

    # This function extracts embedded PDF text when richer local extractors are unavailable.
    def _extract_with_pypdf(self, source_path: str | Path) -> ExtractionResult:
        reader = PdfReader(str(source_path))
        pages: list[ExtractedPage] = []
        warnings: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = normalize_text(page.extract_text() or "") or ""
            if not text:
                warnings.append(f"Low-text or image-only PDF page detected: {page_number}")
            pages.append(ExtractedPage(number=page_number, text=text))
        return ExtractionResult(
            extractor_name="pypdf",
            extraction_mode="fallback_library",
            pages=pages,
            warnings=warnings,
        )


class DocxFallbackExtractor:
    """Extracts DOCX paragraphs and tables locally with python-docx."""

    # This function preserves paragraph style signals and table matrices for canonical assembly.
    def extract(self, source_path: str | Path) -> ExtractionResult:
        document = DocxDocument(str(source_path))
        blocks = [
            {
                "text": normalize_text(paragraph.text) or "",
                "style_name": paragraph.style.name if paragraph.style else "",
            }
            for paragraph in document.paragraphs
            if normalize_text(paragraph.text)
        ]
        tables = [
            [[normalize_text(cell.text) or "" for cell in row.cells] for row in table.rows]
            for table in document.tables
        ]
        warnings = [] if blocks else ["No paragraph text detected in DOCX file"]
        return ExtractionResult(
            extractor_name="python-docx",
            extraction_mode="fallback_library",
            blocks=blocks,
            tables=tables,
            warnings=warnings,
        )


class PowerPointFallbackExtractor:
    """Extracts PPTX slide text, tables, notes, and image metadata with python-pptx."""

    # This function preserves slide-level raw content without applying vision analysis.
    def extract(self, source_path: str | Path) -> ExtractionResult:
        presentation = Presentation(str(source_path))
        pages: list[ExtractedPage] = []
        warnings: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            blocks: list[dict[str, Any]] = []
            tables: list[list[list[str]]] = []
            images: list[dict[str, Any]] = []
            for shape in self._walk_shapes(slide.shapes):
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        text = normalize_text(paragraph.text)
                        if text:
                            blocks.append(
                                {
                                    "text": text,
                                    "level": paragraph.level,
                                    "is_title": bool(
                                        getattr(shape, "is_placeholder", False)
                                        and shape.placeholder_format.idx == 0
                                    ),
                                }
                            )
                elif getattr(shape, "has_table", False):
                    tables.append(
                        [[normalize_text(cell.text) or "" for cell in row.cells] for row in shape.table.rows]
                    )
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    images.append(
                        {
                            "shape_name": shape.name,
                            "bbox": {"left": shape.left, "top": shape.top, "width": shape.width, "height": shape.height},
                        }
                    )
            note_text = self._extract_notes(slide)
            if not blocks:
                warnings.append(f"Low-text or visual-only slide detected: {slide_number}")
            pages.append(
                ExtractedPage(
                    number=slide_number,
                    text="\n".join(block["text"] for block in blocks),
                    text_blocks=blocks,
                    tables=tables,
                    images=images,
                    notes=note_text,
                )
            )
        return ExtractionResult(
            extractor_name="python-pptx",
            extraction_mode="fallback_library",
            pages=pages,
            warnings=warnings,
        )

    def _walk_shapes(self, shapes: Any):
        """Includes text and tables nested inside grouped presentation shapes."""
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from self._walk_shapes(shape.shapes)
            else:
                yield shape

    # This function extracts speaker-note text without treating it as slide-body content.
    def _extract_notes(self, slide: Any) -> str | None:
        if not getattr(slide, "has_notes_slide", False):
            return None
        lines = [
            normalize_text(shape.text) or ""
            for shape in slide.notes_slide.shapes
            if getattr(shape, "has_text_frame", False) and normalize_text(shape.text)
        ]
        return normalize_text(" ".join(lines))
