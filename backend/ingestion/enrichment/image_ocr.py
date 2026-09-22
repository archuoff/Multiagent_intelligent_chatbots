"""Local OCR enrichment for image nodes in canonical Master JSON.

This module reads embedded images from supported source files and adds OCR
evidence to canonical image-node attributes. It performs no remote calls; OCR
is local and non-blocking so extraction remains usable when OCR fails.
"""

from __future__ import annotations

import hashlib
import os
import re
from pathlib import Path
from typing import Any, Callable

from backend.ingestion.models import CanonicalDocument
from backend.ingestion.models import CanonicalNode
from backend.ingestion.models import NodeType
from backend.ingestion.models import SourceType
from backend.ingestion.utils import normalize_text


ImageRecord = dict[str, Any]
ImageProvider = Callable[[Path], dict[tuple[Any, ...], ImageRecord]]
OcrFactory = Callable[[], Any]


class ImageOcrEnrichment:
    """Adds local OCR text to image nodes without changing source extraction."""

    def __init__(self, ocr_factory: OcrFactory | None = None, image_provider: ImageProvider | None = None,
                 asset_root: Path | str = Path("storage") / "visual-assets") -> None:
        """Allows tests to inject OCR and image loading while production stays lazy."""
        self._ocr_factory = ocr_factory
        self._image_provider = image_provider
        self._asset_root = Path(asset_root)
        self._ocr = None

    def apply(self, document: CanonicalDocument) -> int:
        """Runs OCR for supported image nodes and returns the processed count."""
        if os.getenv("JLR_IMAGE_OCR", "true").lower() != "true":
            return 0
        image_nodes = [node for root in document.root_nodes for node in self._walk(root) if node.node_type == NodeType.IMAGE]
        if not image_nodes:
            return 0
        if document.source_type == SourceType.XLSX:
            return self._apply_saved_image_ocr(image_nodes)
        if document.source_type == SourceType.PDF:
            return self._apply_pdf_image_ocr(document, image_nodes)
        if document.source_type not in {SourceType.PPT, SourceType.PPTX}:
            return 0
        try:
            images = (self._image_provider or self._pptx_images)(Path(document.source_path))
        except Exception as error:
            for node in image_nodes:
                self._mark_ocr_failure(node, f"Image extraction failed: {type(error).__name__}")
            return 0
        processed = 0
        for node in image_nodes:
            image, match_type = self._match_image(node, images)
            if image is None:
                self._mark_ocr_failure(node, "No matching embedded image found for this image node.")
                continue
            self._persist_image(document, node, image)
            self._apply_ocr(node, bytes(image["blob"]), match_type)
            processed += 1
        return processed

    def _apply_pdf_image_ocr(self, document: CanonicalDocument, image_nodes: list[CanonicalNode]) -> int:
        """Persists PDF image assets or page-region crops, then runs OCR on each stored image."""
        try:
            self._persist_pdf_images(document, image_nodes)
        except Exception as error:
            for node in image_nodes:
                self._mark_ocr_failure(node, f"PDF image persistence failed: {type(error).__name__}")
            return 0
        return self._apply_saved_image_ocr(image_nodes)

    def _persist_pdf_images(self, document: CanonicalDocument, image_nodes: list[CanonicalNode]) -> None:
        """Stores PDF image bytes so OCR/VLM enrichment can use the standard saved_path flow."""
        try:
            import pymupdf as fitz
        except ImportError as error:  # pragma: no cover - depends on optional local dependency.
            raise RuntimeError("PDF image persistence requires PyMuPDF.") from error
        folder = self._asset_folder(document)
        with fitz.open(str(document.source_path)) as pdf:
            for index, node in enumerate(image_nodes, start=1):
                if node.attributes.get("saved_path"):
                    continue
                page_number = node.provenance.page_number or self._page_from_attributes(node)
                if not page_number or page_number < 1 or page_number > len(pdf):
                    self._mark_ocr_failure(node, "PDF image node has no valid page number for asset persistence.")
                    continue
                page = pdf[page_number - 1]
                try:
                    record = self._pdf_embedded_image(pdf, node) or self._pdf_region_crop(page, node)
                except Exception as error:
                    self._mark_ocr_failure(node, f"PDF image asset extraction failed: {type(error).__name__}")
                    continue
                if record is None:
                    self._mark_ocr_failure(node, "PDF image node has neither extractable xref nor usable bbox.")
                    continue
                folder.mkdir(parents=True, exist_ok=True)
                digest = hashlib.sha256(record["blob"]).hexdigest()
                target = folder / f"page_{page_number:03d}_image_{index:03d}_{digest[:12]}.{record['extension']}"
                target.write_bytes(record["blob"])
                node.attributes.update({
                    "saved_path": target.as_posix(),
                    "image_hash": digest,
                    "image_extension": record["extension"],
                    "image_storage_status": record["status"],
                    "image_persistence_method": record["method"],
                })

    def _pdf_embedded_image(self, pdf: Any, node: CanonicalNode) -> ImageRecord | None:
        """Extracts a native PDF image when xref metadata is available."""
        xref = node.attributes.get("xref")
        if xref is None:
            return None
        try:
            extracted = pdf.extract_image(int(xref))
        except Exception:
            return None
        blob = extracted.get("image")
        if not blob:
            return None
        extension = self._safe_extension(str(extracted.get("ext") or node.attributes.get("extension") or "png"))
        return {"blob": bytes(blob), "extension": extension, "status": "stored", "method": "pdf_xref"}

    def _pdf_region_crop(self, page: Any, node: CanonicalNode) -> ImageRecord | None:
        """Renders a PDF image node's bbox when only positional metadata is available."""
        rect = self._pdf_rect_from_node(page, node)
        if rect is None or rect.is_empty or rect.width <= 0 or rect.height <= 0:
            return None
        dpi = int(os.getenv("JLR_PDF_IMAGE_CROP_DPI", "220"))
        import pymupdf as fitz
        pixmap = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), clip=rect, alpha=False)
        return {"blob": pixmap.tobytes("png"), "extension": "png", "status": "stored_crop", "method": "pdf_bbox_crop"}

    def _pdf_rect_from_node(self, page: Any, node: CanonicalNode) -> Any | None:
        """Builds a PyMuPDF rectangle from common bbox shapes used by Docling/fallback extractors."""
        bbox = node.attributes.get("bbox") or node.provenance.bbox
        if not isinstance(bbox, dict):
            return None
        try:
            import pymupdf as fitz
            if all(key in bbox for key in ("x0", "y0", "x1", "y1")):
                rect = fitz.Rect(float(bbox["x0"]), float(bbox["y0"]), float(bbox["x1"]), float(bbox["y1"]))
            elif all(key in bbox for key in ("left", "top", "right", "bottom")):
                rect = fitz.Rect(float(bbox["left"]), float(bbox["top"]), float(bbox["right"]), float(bbox["bottom"]))
            elif all(key in bbox for key in ("left", "top", "width", "height")):
                left = float(bbox["left"])
                top = float(bbox["top"])
                rect = fitz.Rect(left, top, left + float(bbox["width"]), top + float(bbox["height"]))
            else:
                return None
            page_rect = page.rect
            if rect.y0 > page_rect.height or rect.y1 > page_rect.height:
                rect = fitz.Rect(rect.x0, page_rect.height - rect.y1, rect.x1, page_rect.height - rect.y0)
            return rect & page_rect
        except Exception:
            return None

    def _apply_saved_image_ocr(self, image_nodes: list[CanonicalNode]) -> int:
        """Runs OCR for canonical image nodes that already point to stored image files."""
        processed = 0
        for node in image_nodes:
            saved_path = Path(str(node.attributes.get("saved_path") or ""))
            if not saved_path.is_file():
                self._mark_ocr_failure(node, "No stored image file is available for OCR.")
                continue
            try:
                self._apply_ocr(node, saved_path.read_bytes(), "saved_path")
                processed += 1
            except Exception as error:
                self._mark_ocr_failure(node, f"Saved image OCR failed: {type(error).__name__}")
        return processed

    def _apply_ocr(self, node: CanonicalNode, image_bytes: bytes, match_type: str) -> None:
        """Adds OCR fields to one image node while preserving existing metadata."""
        try:
            result = self._ocr_engine()(image_bytes)
            blocks = self._ocr_blocks(result)
            text = "\n".join(block["text"] for block in blocks if block["text"])
            node.attributes.update({
                "ocr_status": "success" if text else "no_text",
                "ocr_engine": "rapidocr",
                "ocr_text": text,
                "ocr_confidence": self._average_confidence(blocks),
                "ocr_blocks": blocks,
                "ocr_match": match_type,
                "ocr_warnings": [],
            })
        except Exception as error:
            self._mark_ocr_failure(node, f"OCR failed: {type(error).__name__}")

    def _ocr_engine(self) -> Any:
        """Initializes RapidOCR only when an OCR-enabled ingestion actually runs."""
        if self._ocr is None:
            if self._ocr_factory is not None:
                self._ocr = self._ocr_factory()
            else:
                from rapidocr import RapidOCR
                self._ocr = RapidOCR()
        return self._ocr

    def _ocr_blocks(self, result: Any) -> list[dict[str, Any]]:
        """Normalizes RapidOCR output into serializable text blocks."""
        texts = list(getattr(result, "txts", []) or [])
        scores = list(getattr(result, "scores", []) or [])
        boxes = getattr(result, "boxes", None)
        min_score = float(os.getenv("JLR_IMAGE_OCR_MIN_CONFIDENCE", "0.5"))
        blocks: list[dict[str, Any]] = []
        for index, raw_text in enumerate(texts):
            text = normalize_text(str(raw_text)) or ""
            confidence = float(scores[index]) if index < len(scores) and scores[index] is not None else None
            if not text or (confidence is not None and confidence < min_score):
                continue
            box = None
            if boxes is not None and index < len(boxes):
                try:
                    box = boxes[index].tolist()
                except AttributeError:
                    box = boxes[index]
            blocks.append({"text": text, "confidence": confidence, "bbox": box})
        return blocks

    def _average_confidence(self, blocks: list[dict[str, Any]]) -> float | None:
        """Calculates a compact confidence score for UI/audit display."""
        scores = [block["confidence"] for block in blocks if isinstance(block.get("confidence"), float)]
        return round(sum(scores) / len(scores), 4) if scores else None

    def _match_image(self, node: CanonicalNode, images: dict[tuple[Any, ...], ImageRecord]) -> tuple[ImageRecord | None, str]:
        """Finds an embedded image by exact shape+bbox first, then bbox-only fallback."""
        slide_number = node.provenance.slide_number
        bbox = self._bbox_tuple(node.attributes.get("bbox"))
        shape_name = normalize_text(str(node.attributes.get("shape_name") or "")) or ""
        full_key = ("pptx-image", slide_number, shape_name, *bbox)
        if full_key in images:
            return images[full_key], "shape_name_bbox"
        bbox_key = ("pptx-image-bbox", slide_number, *bbox)
        return images.get(bbox_key), "bbox"

    def _persist_image(self, document: CanonicalDocument, node: CanonicalNode, image: ImageRecord) -> None:
        """Stores the extracted image bytes and records the relative path in the node."""
        extension = self._safe_extension(str(image.get("extension") or "png"))
        slide_number = node.provenance.slide_number or 0
        shape_name = self._safe_component(str(node.attributes.get("shape_name") or "image"))
        digest = str(image.get("hash") or "")[:12]
        folder = self._asset_folder(document)
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / f"slide_{slide_number:03d}_{shape_name}_{digest}.{extension}"
        target.write_bytes(bytes(image["blob"]))
        node.attributes.update({
            "saved_path": target.as_posix(),
            "image_hash": image.get("hash"),
            "image_extension": extension,
            "image_storage_status": "stored",
        })

    def _asset_folder(self, document: CanonicalDocument) -> Path:
        """Builds the shared visual-asset folder for one document version."""
        return self._asset_root / self._safe_component(document.agent_id) / self._safe_component(document.document_id) / self._safe_component(document.version)

    def _page_from_attributes(self, node: CanonicalNode) -> int | None:
        """Reads page number from image attributes when provenance is incomplete."""
        value = node.attributes.get("page_number")
        return int(value) if isinstance(value, int) or str(value).isdigit() else None

    def _pptx_images(self, source_path: Path) -> dict[tuple[Any, ...], ImageRecord]:
        """Extracts embedded PPTX picture bytes keyed by slide, shape name, and bbox."""
        import hashlib

        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(str(source_path))
        images: dict[tuple[Any, ...], ImageRecord] = {}
        for slide_number, slide in enumerate(presentation.slides, start=1):
            for shape in self._walk_shapes(slide.shapes, MSO_SHAPE_TYPE):
                bbox = self._shape_bbox_tuple(shape)
                name = normalize_text(str(getattr(shape, "name", "") or "")) or ""
                blob = shape.image.blob
                image = {
                    "blob": blob,
                    "extension": self._safe_extension(str(shape.image.ext or "png")),
                    "hash": hashlib.sha256(blob).hexdigest(),
                }
                images[("pptx-image", slide_number, name, *bbox)] = image
                images[("pptx-image-bbox", slide_number, *bbox)] = image
        return images

    def _walk_shapes(self, shapes: Any, shape_types: Any):
        """Yields pictures inside slides and grouped shapes."""
        for shape in shapes:
            if shape.shape_type == shape_types.GROUP:
                yield from self._walk_shapes(shape.shapes, shape_types)
            elif shape.shape_type == shape_types.PICTURE:
                yield shape

    def _shape_bbox_tuple(self, shape: Any) -> tuple[int, int, int, int]:
        """Normalizes a python-pptx shape box to integer EMU coordinates."""
        return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))

    def _bbox_tuple(self, bbox: Any) -> tuple[int, int, int, int]:
        """Normalizes canonical bbox metadata to integer EMU coordinates."""
        if not isinstance(bbox, dict):
            return (0, 0, 0, 0)
        return tuple(int(float(bbox.get(key, 0) or 0)) for key in ("left", "top", "width", "height"))

    def _mark_ocr_failure(self, node: CanonicalNode, warning: str) -> None:
        """Records OCR failure on a node without failing ingestion."""
        node.attributes.update({
            "ocr_status": "failed",
            "ocr_engine": "rapidocr",
            "ocr_text": "",
            "ocr_confidence": None,
            "ocr_blocks": [],
            "ocr_warnings": [warning],
        })

    def _safe_component(self, value: str) -> str:
        """Creates deterministic filesystem-safe names for generated visual assets."""
        cleaned = re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("._")
        return cleaned or "unknown"

    def _safe_extension(self, value: str) -> str:
        """Keeps image extensions simple and path-safe."""
        cleaned = re.sub(r"[^A-Za-z0-9]+", "", value.lower())
        return cleaned or "png"

    def _walk(self, node: CanonicalNode):
        """Yields the canonical tree in source order."""
        yield node
        for child in node.children:
            yield from self._walk(child)
