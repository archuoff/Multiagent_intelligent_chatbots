"""Compare saved canonical content with native source text before chunking.

Reports token coverage as a diagnostic proxy, not an extraction accuracy score.
It does not measure OCR, images, reading order, table alignment, or correctness.
"""

import argparse
from collections import Counter
import json
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET


def xml_text(root: ET.Element) -> str:
    """Joins Office text runs within paragraphs so formatting cannot split words."""
    paragraphs = []
    for paragraph in root.iter():
        if paragraph.tag.rsplit("}", 1)[-1] != "p":
            continue
        fragments = []
        for element in paragraph.iter():
            local_name = element.tag.rsplit("}", 1)[-1]
            if local_name == "t":
                fragments.append(element.text or "")
            elif local_name in {"tab", "br", "cr"}:
                fragments.append(" ")
        paragraphs.append("".join(fragments))
    return "\n".join(paragraphs)


def source_text(path: Path) -> str:
    """Reads native source text independently of the canonical parser."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, data_only=False, read_only=True)
        try:
            return "\n".join(str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row if value is not None)
        finally:
            workbook.close()
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    with zipfile.ZipFile(path) as archive:
        names = ["word/document.xml"] if path.suffix.lower() == ".docx" else [name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
        return "\n".join(xml_text(ET.fromstring(archive.read(name))) for name in names)


def canonical_text(document: dict) -> str:
    """Reads canonical text and table headers without counting raw backup metadata."""
    parts = []
    def visit(node):
        """Collects one canonical node's readable content and descendants."""
        if node.get("text"):
            parts.append(node["text"])
        for row in node.get("attributes", {}).get("header_rows", []):
            parts.extend(str(value) for value in row if value is not None)
        for child in node.get("children", []):
            visit(child)
    for root in document["root_nodes"]:
        visit(root)
    return "\n".join(parts)


def tokens(text: str) -> Counter:
    """Normalizes words and numbers for a reproducible coverage comparison."""
    return Counter(re.findall(r"\w+", text.casefold()))


def main() -> int:
    """Prints coverage and missing-token examples for a saved canonical artifact."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("canonical_json", type=Path)
    args = parser.parse_args()
    expected = tokens(source_text(args.source))
    actual = tokens(canonical_text(json.loads(args.canonical_json.read_text(encoding="utf-8"))))
    total = sum(expected.values())
    covered = sum((expected & actual).values())
    print(json.dumps({"source_tokens": total, "covered_tokens": covered, "native_text_token_coverage": covered / total if total else None, "missing_token_examples": (expected - actual).most_common(20), "limitation": "Coverage proxy only; source layout, tables and scanned/image content require separate review."}, indent=2, ensure_ascii=True))
    return 0 if total else 1


if __name__ == "__main__":
    raise SystemExit(main())
