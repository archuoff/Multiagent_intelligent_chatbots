"""Regression checks for preservation of source content before chunking."""

import unittest
from types import SimpleNamespace
from openpyxl import Workbook
from unittest.mock import patch
from backend.ingestion.parsers.document import DoclingCanonicalBuilder
from backend.ingestion.parsers.excel import ExcelWorkbookParser
from backend.ingestion.models import NodeType, SourceType


class ContentPreservationTests(unittest.TestCase):
    """Checks exact values and source ordering rather than only node counts."""

    def test_bullet_removal_preserves_words_and_negative_numbers(self):
        """Only explicit bullet prefixes may be removed."""
        builder = DoclingCanonicalBuilder()
        self.assertEqual(builder._strip_bullet("operating pressure"), "operating pressure")
        self.assertEqual(builder._strip_bullet("-20 C"), "-20 C")
        self.assertEqual(builder._strip_bullet("- operating pressure"), "operating pressure")
        self.assertEqual(builder._strip_bullet("-"), "")

    def test_docling_body_order_and_repetition(self):
        """Repeated source paragraphs remain separate and metadata labels stay out of text."""
        exported = {"body": {"children": [{"$ref": "#/texts/1"}, {"$ref": "#/texts/0"}]}, "texts": [{"label": "text", "text": "Repeated"}, {"label": "text", "text": "Repeated"}]}
        nodes = DoclingCanonicalBuilder().build_section_nodes(document_id="doc", source_type=SourceType.DOCX, version="v1", parent_node_id="root", exported=exported, markdown="")
        self.assertEqual([n.text for n in nodes], ["Repeated", "Repeated"])
        self.assertEqual(nodes[0].attributes["source_ref"], "#/texts/1")

    def test_excel_blank_header_and_formula_survive(self):
        """Unlabelled columns and formulas remain available with exact cell coordinates."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sensor", None, "Value"])
        sheet.append(["Radar", "important unlabelled value", "=1+2"])
        sheet.append(["Camera", "None", "real value"])
        parser = ExcelWorkbookParser()
        region = parser._detect_regions(sheet)[0]
        self.assertIn("important unlabelled value", region["rows"][0]["values"].values())
        self.assertIn({"coordinate": "C2", "value": "=1+2"}, region["source_cells"])
        row = parser._build_row_node(document_id="doc", source_type=SourceType.XLSX, version="v1",
            parent_node_id="table", sheet_name="Sheet", region=region, row_data=region["rows"][0])
        self.assertNotIn("None", row.model_dump_json())
        placeholder_row = parser._build_row_node(document_id="doc", source_type=SourceType.XLSX, version="v1",
            parent_node_id="table", sheet_name="Sheet", region=region, row_data=region["rows"][1])
        self.assertNotIn("None", placeholder_row.model_dump_json())
        workbook.close()

    def test_excel_formula_errors_are_flagged_and_excluded_from_semantic_text(self):
        """Source formula errors are preserved for audit but excluded from retrieval text."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sensor", "Status"])
        sheet.append(["Radar", "#VALUE!"])
        sheet["B2"].data_type = "e"
        parser = ExcelWorkbookParser()
        region = parser._detect_regions(sheet)[0]
        row = parser._build_row_node(document_id="doc", source_type=SourceType.XLSX, version="v1",
            parent_node_id="table", sheet_name="Sheet", region=region, row_data=region["rows"][0])
        error_cell = row.children[1]
        self.assertTrue(error_cell.attributes["formula_error"])
        self.assertFalse(error_cell.attributes["retrieval_allowed"])
        self.assertNotIn("#VALUE!", row.attributes["semantic_text"])
        workbook.close()

    def test_excel_images_become_canonical_image_nodes(self):
        """Embedded workbook images are represented as traceable canonical image nodes."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Radar Sensor"
        sheet.append(["Heading"])
        fake_marker = SimpleNamespace(row=57, col=2)
        fake_anchor = SimpleNamespace(_from=fake_marker)
        fake_image = SimpleNamespace(anchor=fake_anchor, width=120, height=80, format="png", path="/xl/media/image1.png")
        fake_image._data = lambda: b"fake-png"
        sheet._images.append(fake_image)
        parser = ExcelWorkbookParser()
        with patch.object(parser, "_persist_excel_image", return_value=SimpleNamespace(as_posix=lambda: "storage/visual-assets/test.png")):
            sheet_node = parser._build_sheet_node(document_id="doc", source_type=SourceType.XLSX, version="v1",
                agent_id="adas-agent", parent_node_id="workbook", sheet_index=0, worksheet=sheet)
        image_nodes = [node for node in sheet_node.children if node.node_type == NodeType.IMAGE]
        self.assertEqual(len(image_nodes), 1)
        self.assertEqual(image_nodes[0].provenance.sheet_name, "Radar Sensor")
        self.assertEqual(image_nodes[0].attributes["anchor_cell"], "C58")
        self.assertEqual(image_nodes[0].attributes["saved_path"], "storage/visual-assets/test.png")
        self.assertEqual(sheet_node.attributes["embedded_image_count"], 1)
        workbook.close()

    def test_excel_short_notes_are_not_truncated(self):
        """Every short narrative row survives beyond the previous three-note limit."""
        workbook = Workbook()
        sheet = workbook.active
        for value in ["one", "two", "three", "four", "five"]:
            sheet.append([value])
        region = ExcelWorkbookParser()._detect_regions(sheet)[0]
        self.assertEqual(region["notes"], ["two", "three", "four", "five"])
        workbook.close()
