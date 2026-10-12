"""Raw XLSX extraction for the new ingestion system.

Openpyxl is authoritative for workbook and worksheet values. The OOXML
package is inspected only for embedded media and drawing relationships because
those relationships are not reliably exposed as a complete public openpyxl
API. This module intentionally does not build the project's canonical schema.
"""

from __future__ import annotations

import hashlib
import json
import mimetypes
import posixpath
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree


class ExcelExtractionError(RuntimeError):
    """Base error for the isolated XLSX extractor."""


class InvalidWorkbookError(ExcelExtractionError):
    """Raised when a source is not a readable XLSX workbook."""


@dataclass(slots=True)
class ExcelCell:
    coordinate: str
    raw_value: Any = None
    formula: str | None = None
    cached_value: Any = None
    data_type: str | None = None
    formula_data_type: str | None = None
    number_format: str | None = None
    hyperlink: dict[str, Any] | None = None
    comment: dict[str, Any] | None = None
    image_formula: dict[str, Any] | None = None


@dataclass(slots=True)
class ExcelTable:
    name: str
    display_name: str | None
    reference: str
    columns: list[str] = field(default_factory=list)
    style: dict[str, Any] | None = None


@dataclass(slots=True)
class ExcelWorksheet:
    name: str
    order: int
    visibility: str
    dimensions: str
    max_row: int
    max_column: int
    cells: list[ExcelCell] = field(default_factory=list)
    merged_ranges: list[str] = field(default_factory=list)
    tables: list[ExcelTable] = field(default_factory=list)


@dataclass(slots=True)
class ExcelImage:
    image_id: str
    status: str
    content_type: str | None = None
    sha256: str | None = None
    worksheet_name: str | None = None
    anchor_start: str | None = None
    anchor_end: str | None = None
    anchors: list[dict[str, Any]] = field(default_factory=list)
    drawing_ids: list[str] = field(default_factory=list)
    provenance: list[str] = field(default_factory=list)
    pixel_width: int | None = None
    pixel_height: int | None = None
    asset_reference: str | None = None
    formula: str | None = None
    source_cell: str | None = None
    external_url: str | None = None
    binary_data: bytes | None = field(default=None, repr=False)


@dataclass(slots=True)
class ExcelExtractionResult:
    """Raw, typed XLSX extraction result."""

    source_metadata: dict[str, Any]
    workbook_metadata: dict[str, Any]
    worksheets: list[ExcelWorksheet]
    images: list[ExcelImage]
    warnings: list[str]
    errors: list[str]
    extraction_configuration: dict[str, Any]
    source_path: Path = field(repr=False)

    def to_dict(self) -> dict[str, Any]:
        """Return a JSON-serializable raw export, excluding in-memory bytes."""
        payload = _json_safe(asdict(self))
        payload.pop("source_path", None)
        return payload

    def to_json(self, *, indent: int = 2) -> str:
        """Serialize the raw export without embedding binary image bytes."""
        return json.dumps(self.to_dict(), indent=indent, sort_keys=True)


_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_DRAWING_MAIN_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_IMAGE_FORMULA_RE = re.compile(r"IMAGE\s*\(", re.IGNORECASE)
_IMAGE_LITERAL_RE = re.compile(r'IMAGE\s*\(\s*"([^"]*)"', re.IGNORECASE)
_IMAGE_REFERENCE_RE = re.compile(
    r"IMAGE\s*\(\s*(?:(?:'([^']+)'|([A-Za-z_][\w ]*))!)?\$?([A-Z]{1,3})\$?(\d+)",
    re.IGNORECASE,
)
_MEDIA_EXTENSIONS = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".gif": "image/gif", ".bmp": "image/bmp", ".tif": "image/tiff",
    ".tiff": "image/tiff", ".webp": "image/webp",
}


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_safe(item) for item in value]
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, bytes):
        return None
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _relationship_target(source: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source), target))


def _read_relationships(archive: ZipFile, source: str) -> dict[str, str]:
    rels_path = posixpath.join(
        posixpath.dirname(source), "_rels", f"{posixpath.basename(source)}.rels"
    )
    try:
        root = ElementTree.fromstring(archive.read(rels_path))
    except KeyError:
        return {}
    return {
        element.attrib["Id"]: _relationship_target(source, element.attrib["Target"])
        for element in root.findall(f"{{{_REL_NS}}}Relationship")
        if "Id" in element.attrib and "Target" in element.attrib
    }


def _cell_coordinate(column: int, row: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row}"


def _anchor_cell(anchor: Any, attribute: str) -> str | None:
    point = getattr(anchor, attribute, None)
    if point is None:
        return None
    return _cell_coordinate(int(point.col) + 1, int(point.row) + 1)


def _openpyxl_anchor(image: Any) -> dict[str, Any]:
    anchor = getattr(image, "anchor", None)
    start = _anchor_cell(anchor, "_from")
    end = _anchor_cell(anchor, "to")
    result: dict[str, Any] = {}
    if start:
        result["start"] = start
    if end:
        result["end"] = end
    return result


def _image_content_type(image: Any) -> str | None:
    content_type = getattr(image, "content_type", None)
    if content_type:
        return content_type
    guessed, _ = mimetypes.guess_type(getattr(image, "path", ""))
    return guessed or None


def _defined_names(workbook: Any) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    collection = getattr(workbook, "defined_names", {})
    items = collection.items() if hasattr(collection, "items") else (
        (item.name, item) for item in collection
    )
    for name, defined in items:
        result.append({
            "name": name,
            "value": getattr(defined, "attr_text", None),
            "local_sheet_id": getattr(defined, "localSheetId", None),
            "hidden": getattr(defined, "hidden", None),
            "comment": getattr(defined, "comment", None),
        })
    return result


def _workbook_metadata(workbook: Any) -> dict[str, Any]:
    properties = getattr(workbook, "properties", None)
    property_names = (
        "creator", "title", "description", "subject", "keywords", "category",
        "language", "lastModifiedBy", "created", "modified", "lastPrinted",
        "revision", "version", "contentStatus", "identifier", "contentType",
    )
    return {
        "sheet_names": list(workbook.sheetnames),
        "sheets": [
            {"name": sheet.title, "order": index, "visibility": sheet.sheet_state}
            for index, sheet in enumerate(workbook.worksheets)
        ],
        "properties": {
            name: getattr(properties, name, None)
            for name in property_names
            if properties is not None and hasattr(properties, name)
        },
        "defined_names": _defined_names(workbook),
        "active_sheet": getattr(getattr(workbook, "active", None), "title", None),
    }


def _cell_hyperlink(cell: Any) -> dict[str, Any] | None:
    hyperlink = getattr(cell, "hyperlink", None)
    if hyperlink is None:
        return None
    return {
        "target": getattr(hyperlink, "target", None),
        "location": getattr(hyperlink, "location", None),
        "display": getattr(hyperlink, "display", None),
        "tooltip": getattr(hyperlink, "tooltip", None),
    }


def _cell_comment(cell: Any) -> dict[str, Any] | None:
    comment = getattr(cell, "comment", None)
    return None if comment is None else {"text": comment.text, "author": comment.author}


def _image_formula_details(formula: str) -> dict[str, Any]:
    details: dict[str, Any] = {"formula": formula}
    literal = _IMAGE_LITERAL_RE.search(formula)
    if literal:
        details.update({"source_type": "literal_url", "external_url": literal.group(1)})
        details["status"] = "formula_image_external_unresolved"
        return details
    reference = _IMAGE_REFERENCE_RE.search(formula)
    if reference:
        details.update({
            "source_type": "source_cell",
            "source_cell": f"{reference.group(3).upper()}{reference.group(4)}",
        })
        if reference.group(1) or reference.group(2):
            details["source_sheet"] = reference.group(1) or reference.group(2)
        details["status"] = "formula_image_external_unresolved"
        return details
    details["status"] = "formula_image_invalid"
    return details


def _iter_existing_cells(sheet: Any) -> Iterable[Any]:
    """Yield existing openpyxl cells without materializing an empty matrix."""
    existing = getattr(sheet, "_cells", None)
    if isinstance(existing, dict):
        yield from sorted(existing.values(), key=lambda cell: (cell.row, cell.column))
        return
    for row in sheet.iter_rows():
        yield from row


def _worksheet_data(formula_sheet: Any, cached_sheet: Any, order: int) -> tuple[ExcelWorksheet, list[str]]:
    cells: list[ExcelCell] = []
    warnings: list[str] = []
    cached_cells = getattr(cached_sheet, "_cells", {})
    for cell in _iter_existing_cells(formula_sheet):
        formula = cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None
        if formula is None and cell.value is None:
            continue
        cached_cell = cached_cells.get((cell.row, cell.column))
        cached_value = getattr(cached_cell, "value", None) if formula else None
        image_formula = None
        if formula and _IMAGE_FORMULA_RE.search(formula):
            image_formula = _image_formula_details(formula)
            warnings.append(
                f"{formula_sheet.title}!{cell.coordinate}: "
                + ("external IMAGE() source was not downloaded."
                   if image_formula["status"] == "formula_image_external_unresolved"
                   else "IMAGE() formula could not be resolved.")
            )
        cells.append(ExcelCell(
            coordinate=cell.coordinate,
            raw_value=cell.value,
            formula=formula,
            cached_value=cached_value,
            data_type=(getattr(cached_cell, "data_type", None)
                       if formula and cached_cell is not None else cell.data_type),
            formula_data_type=cell.data_type if formula else None,
            number_format=cell.number_format,
            hyperlink=_cell_hyperlink(cell),
            comment=_cell_comment(cell),
            image_formula=image_formula,
        ))
    tables = []
    for table in formula_sheet.tables.values():
        tables.append(ExcelTable(
            name=table.name,
            display_name=getattr(table, "displayName", None),
            reference=table.ref,
            columns=[column.name for column in (table.tableColumns or [])],
            style=(table.tableStyleInfo.__dict__.copy() if table.tableStyleInfo else None),
        ))
    return ExcelWorksheet(
        name=formula_sheet.title,
        order=order,
        visibility=formula_sheet.sheet_state,
        dimensions=formula_sheet.calculate_dimension(),
        max_row=formula_sheet.max_row,
        max_column=formula_sheet.max_column,
        cells=cells,
        merged_ranges=[str(item) for item in formula_sheet.merged_cells.ranges],
        tables=tables,
    ), warnings


def _merge_image(image_map: dict[str, ExcelImage], image: ExcelImage) -> ExcelImage:
    key = image.sha256 or image.image_id
    existing = image_map.get(key)
    if existing is None:
        image_map[key] = image
        return image
    for value in image.provenance + image.drawing_ids:
        target = existing.provenance if value in image.provenance else existing.drawing_ids
        if value not in target:
            target.append(value)
    if image.worksheet_name and not existing.worksheet_name:
        existing.worksheet_name = image.worksheet_name
    for anchor in image.anchors:
        if anchor not in existing.anchors:
            existing.anchors.append(anchor)
    existing.anchor_start = existing.anchor_start or image.anchor_start
    existing.anchor_end = existing.anchor_end or image.anchor_end
    existing.pixel_width = existing.pixel_width or image.pixel_width
    existing.pixel_height = existing.pixel_height or image.pixel_height
    return existing


def _save_asset(image: ExcelImage, artifact_dir: Path) -> None:
    if not image.binary_data or not image.sha256:
        return
    extension = mimetypes.guess_extension(image.content_type or "") or ".bin"
    filename = f"image-{image.sha256[:16]}{extension}"
    image_dir = artifact_dir / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    (image_dir / filename).write_bytes(image.binary_data)
    image.asset_reference = f"images/{filename}"


def _openpyxl_images(workbook: Any, image_map: dict[str, ExcelImage], errors: list[str]) -> None:
    for sheet in workbook.worksheets:
        for image in getattr(sheet, "_images", []) or []:
            try:
                data = image._data()
                digest = hashlib.sha256(data).hexdigest()
                anchor = _openpyxl_anchor(image)
                _merge_image(image_map, ExcelImage(
                    image_id=f"sha256:{digest}",
                    status="embedded_extracted",
                    content_type=_image_content_type(image),
                    sha256=digest,
                    worksheet_name=sheet.title,
                    anchor_start=anchor.get("start"),
                    anchor_end=anchor.get("end"),
                    anchors=[anchor] if anchor else [],
                    provenance=["openpyxl"],
                    pixel_width=getattr(image, "width", None),
                    pixel_height=getattr(image, "height", None),
                    binary_data=data,
                ))
            except Exception as error:
                errors.append(f"{sheet.title}: openpyxl image extraction failed: {error}")


def _package_images(archive: ZipFile, image_map: dict[str, ExcelImage], errors: list[str]) -> None:
    names = set(archive.namelist())
    media_by_path: dict[str, ExcelImage] = {}
    for media_path in sorted(name for name in names if name.startswith("xl/media/")):
        try:
            data = archive.read(media_path)
            digest = hashlib.sha256(data).hexdigest()
            record = ExcelImage(
                image_id=f"sha256:{digest}",
                status="embedded_extracted",
                content_type=_MEDIA_EXTENSIONS.get(Path(media_path).suffix.lower()) or mimetypes.guess_type(media_path)[0],
                sha256=digest,
                provenance=["package_inspection"],
                binary_data=data,
            )
            media_by_path[media_path] = _merge_image(image_map, record)
        except Exception as error:
            errors.append(f"{media_path}: package media extraction failed: {error}")

    workbook_path = "xl/workbook.xml"
    try:
        workbook_root = ElementTree.fromstring(archive.read(workbook_path))
        workbook_rels = _read_relationships(archive, workbook_path)
        sheets: dict[str, str] = {}
        for sheet in workbook_root.findall(f"{{{_MAIN_NS}}}sheets/{{{_MAIN_NS}}}sheet"):
            rel_id = sheet.attrib.get(f"{{{_DOC_REL_NS}}}id")
            target = workbook_rels.get(rel_id or "")
            if target:
                sheets[target] = sheet.attrib.get("name", "")
    except (KeyError, ElementTree.ParseError) as error:
        errors.append(f"OOXML worksheet relationship inspection failed: {error}")
        return

    for worksheet_path, sheet_name in sheets.items():
        try:
            worksheet_root = ElementTree.fromstring(archive.read(worksheet_path))
        except (KeyError, ElementTree.ParseError) as error:
            errors.append(f"{worksheet_path}: worksheet relationship inspection failed: {error}")
            continue
        worksheet_rels = _read_relationships(archive, worksheet_path)
        for drawing in worksheet_root.findall(f"{{{_MAIN_NS}}}drawing"):
            drawing_path = worksheet_rels.get(drawing.attrib.get(f"{{{_DOC_REL_NS}}}id", ""))
            if not drawing_path:
                continue
            try:
                drawing_root = ElementTree.fromstring(archive.read(drawing_path))
            except (KeyError, ElementTree.ParseError) as error:
                errors.append(f"{drawing_path}: drawing inspection failed: {error}")
                continue
            drawing_rels = _read_relationships(archive, drawing_path)
            for anchor in drawing_root:
                if _local_name(anchor.tag) not in {"twoCellAnchor", "oneCellAnchor", "absoluteAnchor"}:
                    continue
                start = _xml_anchor(anchor.find(f"{{{_DRAWING_NS}}}from"))
                end = _xml_anchor(anchor.find(f"{{{_DRAWING_NS}}}to"))
                for blip in anchor.findall(f".//{{{_DRAWING_MAIN_NS}}}blip"):
                    media_path = drawing_rels.get(blip.attrib.get(f"{{{_DOC_REL_NS}}}embed", ""))
                    record = media_by_path.get(media_path or "")
                    if record is None:
                        continue
                    record.worksheet_name = record.worksheet_name or sheet_name
                    record.anchor_start = record.anchor_start or start
                    record.anchor_end = record.anchor_end or end
                    anchor_info = {"worksheet": sheet_name, "start": start, "end": end}
                    if anchor_info not in record.anchors:
                        record.anchors.append(anchor_info)
                    drawing_id = _drawing_id(anchor)
                    if drawing_id and drawing_id not in record.drawing_ids:
                        record.drawing_ids.append(drawing_id)


def _xml_anchor(node: Any) -> str | None:
    if node is None:
        return None
    values = {_local_name(child.tag): child.text for child in node}
    try:
        return _cell_coordinate(int(values["col"]) + 1, int(values["row"]) + 1)
    except (KeyError, TypeError, ValueError):
        return None


def _drawing_id(anchor: Any) -> str | None:
    for node in anchor.findall(f".//{{{_DRAWING_NS}}}cNvPr"):
        return node.attrib.get("id") or node.attrib.get("name")
    return None


def _formula_images(worksheets: list[ExcelWorksheet], images: list[ExcelImage]) -> None:
    for sheet in worksheets:
        for cell in sheet.cells:
            if cell.image_formula:
                details = cell.image_formula
                images.append(ExcelImage(
                    image_id=f"formula:{sheet.name}!{cell.coordinate}",
                    status=details["status"],
                    worksheet_name=sheet.name,
                    provenance=["openpyxl_formula"],
                    formula=cell.formula,
                    source_cell=details.get("source_cell"),
                    external_url=details.get("external_url"),
                ))


def extract_xlsx(source_path: str | Path, *, artifact_dir: str | Path | None = None) -> ExcelExtractionResult:
    """Extract one XLSX workbook using openpyxl plus automatic media inspection."""
    path = Path(source_path).expanduser()
    if not path.exists():
        raise InvalidWorkbookError(f"Workbook does not exist: {path}")
    if not path.is_file():
        raise InvalidWorkbookError(f"Workbook path is not a regular file: {path}")
    if path.suffix.lower() != ".xlsx":
        raise InvalidWorkbookError(f"Workbook must have a .xlsx extension: {path}")
    try:
        with ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise InvalidWorkbookError(f"Workbook contains a corrupt ZIP member: {path}")
            names = set(archive.namelist())
            missing = sorted({"[Content_Types].xml", "xl/workbook.xml"} - names)
            if missing:
                raise InvalidWorkbookError(f"Workbook is missing OOXML members: {', '.join(missing)}")
            ElementTree.fromstring(archive.read("[Content_Types].xml"))
            ElementTree.fromstring(archive.read("xl/workbook.xml"))
    except InvalidWorkbookError:
        raise
    except (BadZipFile, OSError, KeyError, ElementTree.ParseError) as error:
        raise InvalidWorkbookError(f"Workbook is not a readable XLSX package: {path}: {error}") from error

    try:
        import openpyxl
        from openpyxl import load_workbook
    except ImportError as error:
        raise ExcelExtractionError("openpyxl is required for XLSX extraction") from error
    try:
        formula_workbook = load_workbook(path, data_only=False, read_only=False)
        cached_workbook = load_workbook(path, data_only=True, read_only=False)
    except Exception as error:
        raise InvalidWorkbookError(f"openpyxl could not read the XLSX workbook: {path}: {error}") from error

    warnings: list[str] = []
    errors: list[str] = []
    worksheets: list[ExcelWorksheet] = []
    for order, (formula_sheet, cached_sheet) in enumerate(zip(formula_workbook.worksheets, cached_workbook.worksheets)):
        worksheet, sheet_warnings = _worksheet_data(formula_sheet, cached_sheet, order)
        worksheets.append(worksheet)
        warnings.extend(sheet_warnings)
    image_map: dict[str, ExcelImage] = {}
    _openpyxl_images(formula_workbook, image_map, errors)
    with ZipFile(path) as archive:
        _package_images(archive, image_map, errors)
    images = list(image_map.values())
    _formula_images(worksheets, images)
    output_artifact_dir = Path(artifact_dir).expanduser() if artifact_dir else Path("xlsx_artifacts") / path.stem
    for image in images:
        _save_asset(image, output_artifact_dir)

    configuration = {
        "primary_library": "openpyxl",
        "openpyxl_version": getattr(openpyxl, "__version__", None),
        "data_only_loads": [False, True],
        "package_media_inspection": True,
        "package_media_paths": ["xl/media/", "xl/drawings/", "xl/drawings/_rels/", "xl/worksheets/_rels/"],
        "asset_directory": "images",
        "external_image_downloads": False,
        "sparse_cells_only": True,
    }
    return ExcelExtractionResult(
        source_metadata={"filename": path.name, "extension": path.suffix.lower(), "size_bytes": path.stat().st_size, "format": "xlsx"},
        workbook_metadata=_workbook_metadata(formula_workbook),
        worksheets=worksheets,
        images=images,
        warnings=warnings,
        errors=errors,
        extraction_configuration=configuration,
        source_path=path.resolve(),
    )


extract_excel = extract_xlsx


__all__ = [
    "ExcelCell", "ExcelExtractionError", "ExcelExtractionResult", "ExcelImage",
    "ExcelTable", "ExcelWorksheet", "InvalidWorkbookError", "extract_excel", "extract_xlsx",
]
