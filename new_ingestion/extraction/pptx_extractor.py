"""Raw PPTX extraction using python-pptx only.

This module preserves presentation structure and provenance without building a
canonical document schema. It intentionally does not inspect OOXML package
relationships; standard embedded pictures are read through python-pptx.
"""

from __future__ import annotations

import hashlib
import json
import mimetypes
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any


class PptxExtractionError(RuntimeError):
    """Base error for the isolated PPTX extractor."""


class InvalidPresentationError(PptxExtractionError):
    """Raised when the source is not a readable PPTX presentation."""


@dataclass(slots=True)
class PptxRun:
    text: str
    font: dict[str, Any] = field(default_factory=dict)
    hyperlink: dict[str, Any] | None = None


@dataclass(slots=True)
class PptxParagraph:
    text: str
    level: int | None = None
    alignment: Any = None
    runs: list[PptxRun] = field(default_factory=list)


@dataclass(slots=True)
class PptxTextFrame:
    text: str
    paragraphs: list[PptxParagraph] = field(default_factory=list)


@dataclass(slots=True)
class PptxTableCell:
    row: int
    column: int
    text: str
    fill: dict[str, Any] = field(default_factory=dict)
    margins: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class PptxTable:
    rows: int
    columns: int
    left: int
    top: int
    width: int
    height: int
    cells: list[PptxTableCell] = field(default_factory=list)


@dataclass(slots=True)
class PptxShape:
    shape_id: int | None
    name: str | None
    shape_type: str
    order_index: int
    z_order: int
    left: int | None
    top: int | None
    width: int | None
    height: int | None
    rotation: float | None
    placeholder: dict[str, Any] | None = None
    parent_shape_id: int | None = None
    text_frame: PptxTextFrame | None = None
    table: PptxTable | None = None


@dataclass(slots=True)
class PptxImage:
    image_id: str
    slide_number: int
    shape_id: int | None
    filename: str | None
    content_type: str | None
    width: int | None
    height: int | None
    left: int | None
    top: int | None
    sha256: str
    asset_reference: str
    placements: list[dict[str, Any]] = field(default_factory=list)
    binary_data: bytes | None = field(default=None, repr=False)


@dataclass(slots=True)
class PptxSlide:
    slide_number: int
    order: int
    layout_name: str | None
    background: dict[str, Any] | None
    shapes: list[PptxShape] = field(default_factory=list)
    speaker_notes: str | None = None


@dataclass(slots=True)
class PptxExtractionResult:
    """Typed raw PPTX output; not the application's canonical JSON."""

    source_metadata: dict[str, Any]
    presentation_metadata: dict[str, Any]
    slides: list[PptxSlide]
    images: list[PptxImage]
    warnings: list[str]
    errors: list[str]
    extraction_configuration: dict[str, Any]
    source_path: Path = field(repr=False)

    def to_dict(self) -> dict[str, Any]:
        """Return JSON-safe raw output without embedding image bytes."""
        payload = _json_safe(asdict(self))
        payload.pop("source_path", None)
        for image in payload.get("images", []):
            image.pop("binary_data", None)
        return payload

    def to_json(self, *, indent: int = 2) -> str:
        return json.dumps(self.to_dict(), indent=indent, sort_keys=True)


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_safe(item) for item in value]
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, bytes):
        return None
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _fill_info(fill: Any) -> dict[str, Any]:
    result: dict[str, Any] = {}
    try:
        result["type"] = str(fill.type) if fill.type is not None else None
    except Exception:
        return result
    try:
        color = fill.fore_color
        result["color"] = {
            "rgb": str(color.rgb) if color.rgb is not None else None,
            "theme": color.theme_color.name if color.theme_color is not None else None,
            "brightness": color.brightness,
        }
    except Exception:
        pass
    return result


def _font_info(font: Any) -> dict[str, Any]:
    if font is None:
        return {}
    result = {}
    for name in ("name", "size", "bold", "italic", "underline", "strike"):
        try:
            value = getattr(font, name)
            result[name] = float(value.pt) if name == "size" and value is not None else value
        except Exception:
            pass
    try:
        color = font.color
        result["color"] = {
            "rgb": str(color.rgb) if color.rgb is not None else None,
            "theme": color.theme_color.name if color.theme_color is not None else None,
        }
    except Exception:
        pass
    return result


def _hyperlink_info(hyperlink: Any) -> dict[str, Any] | None:
    if hyperlink is None:
        return None
    return {
        "address": getattr(hyperlink, "address", None),
        "sub_address": getattr(hyperlink, "sub_address", None),
        "screen_tip": getattr(hyperlink, "screen_tip", None),
        "text": getattr(hyperlink, "text", None),
    }


def _run_info(run: Any) -> PptxRun:
    hyperlink = _hyperlink_info(getattr(run, "hyperlink", None))
    return PptxRun(text=run.text or "", font=_font_info(getattr(run, "font", None)), hyperlink=hyperlink)


def _text_frame_info(text_frame: Any) -> PptxTextFrame:
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        paragraphs.append(PptxParagraph(
            text=paragraph.text or "",
            level=getattr(paragraph, "level", None),
            alignment=str(getattr(paragraph, "alignment", None)) if getattr(paragraph, "alignment", None) is not None else None,
            runs=[_run_info(run) for run in paragraph.runs],
        ))
    return PptxTextFrame(text=text_frame.text or "", paragraphs=paragraphs)


def _shape_type_name(shape: Any) -> str:
    value = getattr(shape, "shape_type", None)
    return getattr(value, "name", None) or str(value)


def _placeholder_info(shape: Any) -> dict[str, Any] | None:
    try:
        if not shape.is_placeholder:
            return None
        placeholder = shape.placeholder_format
        return {
            "is_placeholder": True,
            "idx": getattr(placeholder, "idx", None),
            "type": str(getattr(placeholder, "type", None)),
        }
    except Exception:
        return {"is_placeholder": True}


def _table_info(shape: Any) -> PptxTable:
    table = shape.table
    cells = []
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cells.append(PptxTableCell(
                row=row_index,
                column=column_index,
                text=cell.text or "",
                fill=_fill_info(cell.fill),
                margins={
                    name: getattr(cell, name, None)
                    for name in ("margin_left", "margin_right", "margin_top", "margin_bottom")
                },
            ))
    return PptxTable(
        rows=len(table.rows),
        columns=len(table.columns),
        left=int(shape.left),
        top=int(shape.top),
        width=int(shape.width),
        height=int(shape.height),
        cells=cells,
    )


def _slide_background(slide: Any) -> dict[str, Any] | None:
    try:
        return _fill_info(slide.background.fill)
    except Exception:
        return None


def _notes_text(slide: Any, warnings: list[str], slide_number: int) -> str | None:
    try:
        notes_slide = slide.notes_slide
        text_frame = getattr(notes_slide, "notes_text_frame", None)
        return text_frame.text if text_frame is not None else None
    except Exception as error:
        warnings.append(f"Slide {slide_number}: speaker notes could not be read: {error}")
        return None


def _warn_for_shape(shape: Any, slide_number: int, warnings: list[str]) -> None:
    shape_type = _shape_type_name(shape).lower()
    if "chart" in shape_type:
        warnings.append(f"Slide {slide_number}, shape {getattr(shape, 'shape_id', '?')}: chart extraction is not implemented.")
    elif "smartart" in shape_type:
        warnings.append(f"Slide {slide_number}, shape {getattr(shape, 'shape_id', '?')}: SmartArt is not supported by python-pptx.")
    elif "ole" in shape_type or "embedded" in shape_type:
        warnings.append(f"Slide {slide_number}, shape {getattr(shape, 'shape_id', '?')}: embedded object extraction is not implemented.")
    elif "media" in shape_type:
        warnings.append(f"Slide {slide_number}, shape {getattr(shape, 'shape_id', '?')}: linked or media extraction is unsupported.")
    elif "graphic_frame" in shape_type and not getattr(shape, "has_table", False):
        warnings.append(f"Slide {slide_number}, shape {getattr(shape, 'shape_id', '?')}: SmartArt or unsupported graphic frame is not extracted.")


def _save_image(image: PptxImage, artifact_dir: Path) -> None:
    target = artifact_dir / "images"
    target.mkdir(parents=True, exist_ok=True)
    suffix = mimetypes.guess_extension(image.content_type or "") or Path(image.filename or "").suffix or ".bin"
    filename = f"image-{image.sha256}{suffix}"
    (target / filename).write_bytes(image.binary_data or b"")
    image.asset_reference = f"images/{filename}"


def _extract_shapes(
    shapes: Any,
    slide_number: int,
    warnings: list[str],
    errors: list[str],
    images: dict[str, PptxImage],
    artifact_dir: Path,
    output_shapes: list[PptxShape],
    parent_shape_id: int | None = None,
) -> None:
    for shape in shapes:
        order_index = len(output_shapes)
        shape_id = getattr(shape, "shape_id", None)
        shape_type = _shape_type_name(shape)
        text_frame = None
        table = None
        try:
            if getattr(shape, "has_text_frame", False):
                text_frame = _text_frame_info(shape.text_frame)
            if getattr(shape, "has_table", False):
                table = _table_info(shape)
        except Exception as error:
            errors.append(f"Slide {slide_number}, shape {shape_id}: shape content extraction failed: {error}")
        output_shapes.append(PptxShape(
            shape_id=shape_id,
            name=getattr(shape, "name", None),
            shape_type=shape_type,
            order_index=order_index,
            z_order=order_index,
            left=int(shape.left) if getattr(shape, "left", None) is not None else None,
            top=int(shape.top) if getattr(shape, "top", None) is not None else None,
            width=int(shape.width) if getattr(shape, "width", None) is not None else None,
            height=int(shape.height) if getattr(shape, "height", None) is not None else None,
            rotation=getattr(shape, "rotation", None),
            placeholder=_placeholder_info(shape),
            parent_shape_id=parent_shape_id,
            text_frame=text_frame,
            table=table,
        ))
        _warn_for_shape(shape, slide_number, warnings)
        try:
            if getattr(shape, "shape_type", None) is not None and hasattr(shape, "image"):
                image = shape.image
                data = image.blob
                digest = hashlib.sha256(data).hexdigest()
                placement = {
                    "slide_number": slide_number,
                    "shape_id": shape_id,
                    "left": int(shape.left),
                    "top": int(shape.top),
                    "width": int(shape.width),
                    "height": int(shape.height),
                }
                record = images.get(digest)
                if record is None:
                    record = PptxImage(
                        image_id=f"sha256:{digest}",
                        slide_number=slide_number,
                        shape_id=shape_id,
                        filename=image.filename,
                        content_type=image.content_type,
                        width=image.size[0],
                        height=image.size[1],
                        left=int(shape.left),
                        top=int(shape.top),
                        sha256=digest,
                        asset_reference=f"images/image-{digest}{Path(image.filename or '').suffix or '.bin'}",
                        placements=[placement],
                        binary_data=data,
                    )
                    images[digest] = record
                    _save_image(record, artifact_dir)
                elif placement not in record.placements:
                    record.placements.append(placement)
        except Exception as error:
            if "picture" in shape_type.lower() or hasattr(shape, "image"):
                warnings.append(f"Slide {slide_number}, shape {shape_id}: linked or unextractable media: {error}")
        if hasattr(shape, "shapes"):
            _extract_shapes(shape.shapes, slide_number, warnings, errors, images, artifact_dir, output_shapes, shape_id)


def extract_pptx(source_path: str | Path, artifact_dir: str | Path | None = None) -> PptxExtractionResult:
    """Extract one PPTX presentation using python-pptx only."""
    path = Path(source_path).expanduser()
    if not path.exists():
        raise InvalidPresentationError(f"Presentation does not exist: {path}")
    if not path.is_file():
        raise InvalidPresentationError(f"Presentation path is not a regular file: {path}")
    if path.suffix.lower() != ".pptx":
        raise InvalidPresentationError(f"Presentation must have a .pptx extension: {path}")
    try:
        from pptx import Presentation
    except ImportError as error:
        raise PptxExtractionError("python-pptx is required for PPTX extraction") from error
    try:
        presentation = Presentation(str(path))
    except Exception as error:
        raise InvalidPresentationError(f"python-pptx could not read the presentation: {path}: {error}") from error

    output_artifact_dir = Path(artifact_dir).expanduser() if artifact_dir else Path("pptx_artifacts") / path.stem
    warnings: list[str] = []
    errors: list[str] = []
    image_map: dict[str, PptxImage] = {}
    slides: list[PptxSlide] = []
    for index, slide in enumerate(presentation.slides, start=1):
        shapes: list[PptxShape] = []
        _extract_shapes(slide.shapes, index, warnings, errors, image_map, output_artifact_dir, shapes)
        slides.append(PptxSlide(
            slide_number=index,
            order=index - 1,
            layout_name=getattr(getattr(slide, "slide_layout", None), "name", None),
            background=_slide_background(slide),
            shapes=shapes,
            speaker_notes=_notes_text(slide, warnings, index),
        ))
    properties = presentation.core_properties
    property_names = (
        "author", "title", "subject", "keywords", "comments", "last_modified_by",
        "created", "modified", "category", "content_status", "identifier", "language",
        "revision", "version",
    )
    presentation_metadata = {
        "slide_count": len(presentation.slides),
        "slide_width": int(presentation.slide_width),
        "slide_height": int(presentation.slide_height),
        "coordinate_unit": "EMU",
        "properties": {name: getattr(properties, name, None) for name in property_names},
    }
    return PptxExtractionResult(
        source_metadata={"filename": path.name, "extension": path.suffix.lower(), "size_bytes": path.stat().st_size, "format": "pptx"},
        presentation_metadata=presentation_metadata,
        slides=slides,
        images=list(image_map.values()),
        warnings=warnings,
        errors=errors,
        extraction_configuration={
            "primary_library": "python-pptx",
            "python_pptx_version": getattr(__import__("pptx"), "__version__", None),
            "ooxml_media_inspection": False,
            "image_assets": "images/",
            "external_downloads": False,
            "coordinate_unit": "EMU",
        },
        source_path=path.resolve(),
    )


extract_powerpoint = extract_pptx


__all__ = [
    "InvalidPresentationError", "PptxExtractionError", "PptxExtractionResult",
    "extract_powerpoint", "extract_pptx",
]
