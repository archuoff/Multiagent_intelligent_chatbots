"""Focused tests for the raw python-pptx extractor."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from PIL import Image as PillowImage
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from new_ingestion.extraction.pptx_extractor import (
    InvalidPresentationError,
    extract_pptx,
)


class PptxExtractorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.presentation_path = self.root / "sample.pptx"
        self.artifact_dir = self.root / "assets"
        image_path = self.root / "fixture.png"
        PillowImage.new("RGB", (16, 10), (40, 120, 210)).save(image_path)

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        first = presentation.slides.add_slide(presentation.slide_layouts[6])
        text_box = first.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        text_frame = text_box.text_frame
        text_frame.text = "First slide"
        paragraph = text_frame.add_paragraph()
        paragraph.text = "second paragraph"
        run = paragraph.add_run()
        run.text = " styled"
        run.font.bold = True
        run.font.size = Pt(18)
        run.hyperlink.address = "https://example.invalid"
        first.shapes.add_picture(str(image_path), Inches(1), Inches(2), width=Inches(2))
        first.shapes.add_picture(str(image_path), Inches(4), Inches(2), width=Inches(2))

        second = presentation.slides.add_slide(presentation.slide_layouts[6])
        table_shape = second.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
        table_shape.table.cell(0, 0).text = "Header"
        table_shape.table.cell(0, 1).text = "Value"
        table_shape.table.cell(1, 0).text = "Row"
        table_shape.table.cell(1, 1).text = "42"
        chart_data = ChartData()
        chart_data.categories = ["A", "B"]
        chart_data.add_series("Series 1", (1, 2))
        second.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(1), Inches(4), Inches(3), chart_data)
        presentation.save(self.presentation_path)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_slides_text_runs_tables_and_coordinates(self) -> None:
        result = extract_pptx(self.presentation_path, self.artifact_dir)
        self.assertEqual(result.presentation_metadata["slide_count"], 2)
        self.assertEqual([slide.slide_number for slide in result.slides], [1, 2])
        text_shapes = [shape for shape in result.slides[0].shapes if shape.text_frame]
        self.assertEqual(text_shapes[0].text_frame.paragraphs[0].text, "First slide")
        second_paragraph_runs = text_shapes[0].text_frame.paragraphs[1].runs
        self.assertEqual(second_paragraph_runs[0].text, "second paragraph")
        self.assertEqual(second_paragraph_runs[1].text, " styled")
        self.assertEqual(second_paragraph_runs[1].font["bold"], True)
        self.assertEqual(second_paragraph_runs[1].hyperlink["address"], "https://example.invalid")
        table_shapes = [shape for shape in result.slides[1].shapes if shape.table]
        self.assertEqual(table_shapes[0].table.rows, 2)
        self.assertEqual(table_shapes[0].table.cells[3].text, "42")
        self.assertGreater(text_shapes[0].left, 0)
        self.assertEqual(result.extraction_configuration["ooxml_media_inspection"], False)

    def test_images_are_deduplicated_and_saved_with_placements(self) -> None:
        result = extract_pptx(self.presentation_path, self.artifact_dir)
        self.assertEqual(len(result.images), 1)
        image = result.images[0]
        self.assertEqual(len(image.placements), 2)
        self.assertEqual(image.content_type, "image/png")
        self.assertEqual(image.width, 16)
        self.assertEqual(image.height, 10)
        self.assertTrue(image.asset_reference.startswith("images/image-"))
        asset = self.artifact_dir / image.asset_reference
        self.assertTrue(asset.is_file())
        self.assertEqual(asset.read_bytes(), image.binary_data)

    def test_warnings_and_json_export(self) -> None:
        result = extract_pptx(self.presentation_path, self.artifact_dir)
        self.assertTrue(any("chart extraction" in warning for warning in result.warnings))
        payload = result.to_dict()
        json.dumps(payload)
        self.assertNotIn("source_path", payload)
        self.assertNotIn("binary_data", payload["images"][0])
        self.assertTrue(payload["images"][0]["asset_reference"].startswith("images/"))

    def test_invalid_pptx_input(self) -> None:
        invalid = self.root / "invalid.pptx"
        invalid.write_text("not a presentation", encoding="utf-8")
        with self.assertRaises(InvalidPresentationError):
            extract_pptx(invalid)


if __name__ == "__main__":
    unittest.main()
