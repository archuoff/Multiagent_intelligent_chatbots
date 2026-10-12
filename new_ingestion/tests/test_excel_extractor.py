"""Focused tests for the isolated raw XLSX extractor."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image as PillowImage

from new_ingestion.extraction.excel_extractor import (
    InvalidWorkbookError,
    extract_xlsx,
)


class ExcelExtractorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.workbook_path = self.root / "sample.xlsx"
        self.artifact_dir = self.root / "assets"
        workbook = Workbook()
        visible = workbook.active
        visible.title = "Visible"
        hidden = workbook.create_sheet("Hidden")
        hidden.sheet_state = "hidden"

        visible["A1"] = "text"
        visible["A2"] = 42
        visible["A3"] = True
        visible["A4"] = __import__("datetime").date(2026, 10, 5)
        visible["A5"] = "=1+1"
        visible["A6"] = '=IMAGE("https://example.invalid/image.png")'
        visible["A7"] = "=IMAGE(B1)"
        visible["A8"].comment = Comment("note", "tester")
        visible["A8"].hyperlink = "https://example.invalid"
        visible.merge_cells("C1:D2")
        visible["C1"] = "merged"
        visible["A10"] = "Name"
        visible["B10"] = "Value"
        for row in range(11, 13):
            visible.cell(row, 1, f"row {row}")
            visible.cell(row, 2, row)
        table = Table(displayName="SampleTable", ref="A10:B12")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        visible.add_table(table)
        hidden["A1"] = "not visible"

        image_path = self.root / "fixture.png"
        PillowImage.new("RGB", (12, 8), (30, 90, 180)).save(image_path)
        visible.add_image(OpenpyxlImage(str(image_path)), "F3")
        workbook.save(self.workbook_path)
        self._inject_cached_error(self.workbook_path, "A6", "#VALUE!")

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    @staticmethod
    def _inject_cached_error(path: Path, coordinate: str, value: str) -> None:
        temp_path = path.with_suffix(".rewritten.xlsx")
        with ZipFile(path, "r") as source, ZipFile(temp_path, "w") as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    root = ElementTree.fromstring(data)
                    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                    for cell in root.findall(f".//{{{namespace}}}c"):
                        if cell.attrib.get("r") == coordinate:
                            cell.attrib["t"] = "e"
                            value_node = cell.find(f"{{{namespace}}}v")
                            if value_node is None:
                                value_node = ElementTree.SubElement(cell, f"{{{namespace}}}v")
                            value_node.text = value
                    data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(info, data)
        path.unlink()
        temp_path.rename(path)

    def test_cells_metadata_formulas_cached_values_and_sheet_metadata(self) -> None:
        result = extract_xlsx(self.workbook_path, artifact_dir=self.artifact_dir)
        self.assertEqual(result.workbook_metadata["sheet_names"], ["Visible", "Hidden"])
        self.assertEqual(result.workbook_metadata["sheets"][1]["visibility"], "hidden")
        visible = result.worksheets[0]
        cells = {cell.coordinate: cell for cell in visible.cells}
        self.assertEqual(cells["A1"].raw_value, "text")
        self.assertEqual(cells["A2"].raw_value, 42)
        self.assertIs(cells["A3"].raw_value, True)
        self.assertEqual(cells["A5"].formula, "=1+1")
        self.assertEqual(cells["A6"].cached_value, "#VALUE!")
        self.assertEqual(cells["A6"].data_type, "e")
        self.assertEqual(cells["A6"].formula_data_type, "f")
        self.assertNotIn("A9", cells)
        self.assertEqual(visible.merged_ranges, ["C1:D2"])
        self.assertEqual(visible.tables[0].reference, "A10:B12")
        self.assertEqual(cells["A8"].comment["author"], "tester")
        self.assertEqual(cells["A8"].hyperlink["target"], "https://example.invalid")

    def test_image_formulas_preserve_formula_and_warn_without_download(self) -> None:
        result = extract_xlsx(self.workbook_path, artifact_dir=self.artifact_dir)
        cells = {cell.coordinate: cell for cell in result.worksheets[0].cells}
        self.assertEqual(cells["A6"].formula, '=IMAGE("https://example.invalid/image.png")')
        self.assertEqual(cells["A6"].image_formula["status"], "formula_image_external_unresolved")
        self.assertEqual(cells["A7"].image_formula["source_cell"], "B1")
        formula_images = [image for image in result.images if image.formula]
        self.assertEqual(len(formula_images), 2)
        self.assertTrue(any("external IMAGE() source" in warning for warning in result.warnings))

    def test_openpyxl_and_package_media_are_deduplicated(self) -> None:
        result = extract_xlsx(self.workbook_path, artifact_dir=self.artifact_dir)
        embedded = [image for image in result.images if image.status == "embedded_extracted"]
        self.assertEqual(len(embedded), 1)
        self.assertEqual(embedded[0].provenance, ["openpyxl", "package_inspection"])
        self.assertEqual(embedded[0].worksheet_name, "Visible")
        self.assertEqual(embedded[0].anchor_start, "F3")
        self.assertTrue(embedded[0].asset_reference.startswith("images/image-"))
        self.assertEqual(len(list((self.artifact_dir / "images").glob("*"))), 1)
        self.assertEqual(embedded[0].pixel_width, 12)
        self.assertEqual(embedded[0].pixel_height, 8)

    def test_json_export_is_serializable_and_has_no_absolute_paths(self) -> None:
        result = extract_xlsx(self.workbook_path, artifact_dir=self.artifact_dir)
        payload = result.to_dict()
        json.dumps(payload)
        self.assertNotIn("source_path", payload)
        self.assertEqual(payload["images"][0]["asset_reference"].startswith("images/"), True)

    def test_invalid_xlsx_input_has_extractor_error(self) -> None:
        invalid = self.root / "invalid.xlsx"
        invalid.write_text("not a zip", encoding="utf-8")
        with self.assertRaises(InvalidWorkbookError):
            extract_xlsx(invalid)


if __name__ == "__main__":
    unittest.main()
